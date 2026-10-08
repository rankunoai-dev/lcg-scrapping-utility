"""S8 unit tests — reporter sanitization, prescription gate, validator seeds."""

import json

import openpyxl
import pytest

from src.qa.output_validator import OutputValidator
from src.reporter.report_generator import (FIELD_TASK_LANDMARK, ReportGenerator,
                                           _clean)

PARAMS = json.load(open("config/parameters_database.json", encoding="utf-8"))["parameters"]

REGISTER = {
    "5001": {"school_id": "5001", "locality": "Hunters Creek", "approved": "yes",
             "approved_by": "analyst", "approved_date": "2026-10-06"},
    "5002": {"school_id": "5002", "locality": "Bridle Trails", "approved": ""},
}
NEARBY = {"5001": ["Kirkland"], "5002": ["Kirkland"]}


def param_row(pid, name, action, w="ABSENT", r="ABSENT", pev=None, flags=None,
              priority="High", must=False, cat="GEO"):
    return {
        "parameter_name": name, "category": cat, "priority": priority,
        "must_have": must, "weight": 3.0, "welcome_status": w, "rest_status": r,
        "page_status": w if w != "ABSENT" else r, "location_state": "ABSENT",
        "action_category": action, "welcome_evidence": None, "page_evidence": pev,
        "source_tier": "T", "provenance": "test", "flags": flags or [],
        "counts": {}, "method": "TEST", "confidence": None,
    }


def card(sid, params=None, flags=None):
    return {
        "school_id": sid, "school_name": f"School {sid}", "brand": "B",
        "city": "C", "state": "FL", "s_welcome": 50.0, "s_page": 60.0,
        "s_gap": 10.0, "s_must_have": 80.0, "s_high_priority": 70.0,
        "category_scores_welcome": {}, "mean_confidence": 0.9,
        "priority_band": "Tier 2: Medium Optimization Needed",
        "welcome_found": True, "welcome_word_count": 200,
        "extraction_method": "BRAND_SELECTOR", "welcome_dom_path": "x",
        "flags": flags or [], "parameters": params or {},
        "envelope": {"config_hash": "h", "llm": {}, "fetch_meta": {}},
    }


def build(tmp_path, cards, register=REGISTER, nearby=NEARBY):
    out = tmp_path / "report.xlsx"
    ReportGenerator(out, locality_register=register, nearby_names=nearby) \
        .generate_report(cards, PARAMS, run_context={"config_hash": "h"})
    return out


class TestClean:
    def test_formula_injection_prefixed(self):
        assert _clean("=HYPERLINK(evil)").startswith("'")
        assert _clean("+SUM(A1)").startswith("'")
        assert _clean("@cmd").startswith("'")

    def test_control_chars_stripped(self):
        assert "\x07" not in _clean("ding\x07dong")

    def test_truncation(self):
        assert len(_clean("x" * 40000)) <= 32000

    def test_numbers_untouched(self):
        assert _clean(-0.32) == -0.32


class TestPrescriptionGate:
    def gen(self):
        return ReportGenerator("x.xlsx", locality_register=REGISTER, nearby_names=NEARBY)

    def test_approved_locality_prescribed(self):
        p = param_row("GEO_04", "Locality", "CREATE_NEW_CONTENT")
        act = self.gen().remediation_action("5001", "GEO_04", p)
        assert act is not None and "Hunters Creek" in act[1]
        assert "approved by analyst" in act[2]

    def test_pending_locality_never_in_queue(self):
        p = param_row("GEO_04", "Locality", "CREATE_NEW_CONTENT")
        assert self.gen().remediation_action("5002", "GEO_04", p) is None

    def test_non_register_school_gets_no_locality_task(self):
        p = param_row("GEO_04", "Locality", "CREATE_NEW_CONTENT")
        assert self.gen().remediation_action("9999", "GEO_04", p) is None

    def test_landmark_gap_is_field_task(self):
        p = param_row("GEO_06", "Landmark 1", "CREATE_NEW_CONTENT")
        act = self.gen().remediation_action("5001", "GEO_06", p)
        assert act[0] == "FIELD TASK" and act[1] == FIELD_TASK_LANDMARK


class TestValidatorSeeds:
    def validate(self, path):
        return OutputValidator(REGISTER, NEARBY).validate(path)

    def test_clean_workbook_passes(self, tmp_path):
        p1 = param_row("KW_01", "Daycare", "MOVE_OR_SURFACE", w="ABSENT",
                       r="PRESENT", pev="quality daycare here")
        out = build(tmp_path, [card("5001", {"KW_01": p1})])
        assert self.validate(out) == []

    def test_unapproved_locality_in_queue_caught(self, tmp_path):
        # Sabotage: a reporter whose gate is broken (simulated by editing the
        # finished workbook) must still be caught by the validator.
        out = build(tmp_path, [card("5002")])
        wb = openpyxl.load_workbook(out)
        wb["04_Remediation_Queue"].append(
            ["5002", "School 5002", "http://x", "Tier 2", "GEO_04", "Locality",
             "High", "CREATE NEW CONTENT",
             'ADD locality "Bridle Trails" to the Welcome.', "sabotaged"])
        wb.save(out)
        v = self.validate(out)
        assert any("V3" in x and "Bridle Trails" in x for x in v)

    def test_quoted_page_evidence_is_exempt(self, tmp_path):
        # A MOVE row may QUOTE the school's own page even when the quote
        # contains a register/NEARBY word — that is evidence, not prescription.
        p = param_row("KW_01", "Daycare", "MOVE_OR_SURFACE", w="ABSENT",
                      r="PRESENT", pev="quality daycare here")
        out = build(tmp_path, [card("5002", {"KW_01": p})])
        wb = openpyxl.load_workbook(out)
        wb["04_Remediation_Queue"].append(
            ["5002", "School 5002", "http://x", "Tier 2", "GEO_05", "Street",
             "High", "MOVE / SURFACE",
             "Copy already exists — surface it. Page evidence: our campus at "
             "4580 Bridle Trails Rd near Kirkland city line", "page quote"])
        wb.save(out)
        assert self.validate(out) == []

    def test_nearby_name_in_queue_caught(self, tmp_path):
        out = build(tmp_path, [card("5001")])
        wb = openpyxl.load_workbook(out)
        wb["04_Remediation_Queue"].append(
            ["5001", "School 5001", "http://x", "Tier 2", "GEO_04", "Locality",
             "High", "CREATE NEW CONTENT", "Mention nearby Kirkland families.",
             "sabotaged"])
        wb.save(out)
        v = self.validate(out)
        assert any("V4" in x and "Kirkland" in x for x in v)

    def test_suggested_landmark_caught(self, tmp_path):
        out = build(tmp_path, [card("5001")])
        wb = openpyxl.load_workbook(out)
        wb["04_Remediation_Queue"].append(
            ["5001", "School 5001", "http://x", "Tier 2", "GEO_06", "Landmark 1",
             "High", "CREATE NEW CONTENT", "Say you are across from Starbucks.",
             "sabotaged"])
        wb.save(out)
        v = self.validate(out)
        assert any("V5" in x for x in v)

    def test_injection_text_caught(self, tmp_path):
        out = build(tmp_path, [card("5001")])
        wb = openpyxl.load_workbook(out)
        wb["05_Evaluation_Evidence"].append(
            ["5001", "=HYPERLINK(\"http://evil\")", True, "m", "d", 1, "", "", "", "", 0.9, ""])
        wb.save(out)
        v = self.validate(out)
        assert any("V1" in x for x in v)

    def test_present_without_evidence_caught(self, tmp_path):
        bad = param_row("KW_01", "Daycare", "OPTIMAL", w="PRESENT", r="PRESENT")
        bad["welcome_evidence"] = None
        bad["page_evidence"] = None
        out = build(tmp_path, [card("5001", {"KW_01": bad})])
        v = self.validate(out)
        assert any("V6" in x for x in v)

    def test_reporter_sanitizes_hostile_evidence_end_to_end(self, tmp_path):
        hostile = param_row("KW_01", "Daycare", "MOVE_OR_SURFACE", w="ABSENT",
                            r="PRESENT", pev="=cmd|' /C calc'!A0 \x07bell")
        out = build(tmp_path, [card("5001", {"KW_01": hostile})])
        assert self.validate(out) == []


class TestPresencePivots:
    def make(self, tmp_path):
        p = param_row("KW_01", "Daycare", "KEEP_REINFORCE", w="PRESENT", r="ABSENT")
        p["welcome_evidence"] = "quality daycare in Jacksonville"
        p["page_evidence"] = "quality daycare in Jacksonville"  # welcome fallback
        q = param_row("GEO_04", "Locality", "NOT_APPLICABLE", w="N_A", r="N_A")
        c = card("5001", {"KW_01": p, "GEO_04": q})
        out = tmp_path / "pivot.xlsx"
        ReportGenerator(out, locality_register=REGISTER, nearby_names=NEARBY,
                        school_urls={"5001": "https://example.com/school-5001/"}) \
            .generate_report([c], PARAMS, run_context={"config_hash": "h"})
        return openpyxl.load_workbook(out)

    def test_pivot_sheets_exist_with_urls(self, tmp_path):
        wb = self.make(tmp_path)
        for name in ("07_Welcome_Presence_Pivot", "08_Page_Excl_Welcome_Pivot"):
            ws = wb[name]
            hdr = [c.value for c in ws[1][:3]]
            assert hdr == ["School ID", "School Name", "School URL"]
            assert ws["C2"].value == "https://example.com/school-5001/"

    def test_welcome_present_embeds_quote_and_rest_shows_absent(self, tmp_path):
        wb = self.make(tmp_path)
        pids = [p["id"] for p in PARAMS]
        col = 4 + pids.index("KW_01")
        wel = wb["07_Welcome_Presence_Pivot"].cell(row=2, column=col).value
        rest = wb["08_Page_Excl_Welcome_Pivot"].cell(row=2, column=col).value
        assert wel.startswith("PRESENT — \"quality daycare")
        # whole-page-content sheet does NOT count the welcome section:
        assert rest == "ABSENT"

    def test_na_cells(self, tmp_path):
        wb = self.make(tmp_path)
        pids = [p["id"] for p in PARAMS]
        col = 4 + pids.index("GEO_04")
        assert wb["07_Welcome_Presence_Pivot"].cell(row=2, column=col).value == "N_A"

    def test_pivot_headers_have_hover_tooltips(self, tmp_path):
        wb = self.make(tmp_path)
        ws = wb["07_Welcome_Presence_Pivot"]
        pids = [p["id"] for p in PARAMS]
        col = 4 + pids.index("GEO_06")
        comment = ws.cell(row=1, column=col).comment
        assert comment is not None
        assert "Landmark 1" in comment.text  # hover shows the parameter name

    def test_welcome_present_absent_lists(self, tmp_path):
        wb = self.make(tmp_path)
        ws = wb["09_Welcome_Present_Absent"]
        hdr = [c.value for c in ws[1]]
        assert hdr == ["School ID", "School Name", "School URL", "Welcome Source",
                       "Present", "Partial", "Absent", "Absent on Whole Page"]
        # welcome-absent but page-present (MOVE case) must NOT be in column H
        row0 = next(ws.iter_rows(min_row=2, values_only=True))
        assert "Daycare" not in str(row0[7]) or str(row0[7]) == "(none)"
        row = next(ws.iter_rows(min_row=2, values_only=True))
        present, absent = str(row[4]), str(row[6])
        assert "Daycare" in present                # fixture: PRESENT in welcome
        assert "KW_01" not in present + absent     # names only, no parameter IDs
        assert "Locality" not in present + absent  # N_A stays out of all lists

    def test_every_header_cell_has_hover_meaning(self, tmp_path):
        wb = self.make(tmp_path)
        missing = []
        for ws in wb.worksheets:
            for cell in ws[1]:
                if cell.value is None:
                    continue
                if cell.comment is None:
                    missing.append(f"{ws.title}!{cell.coordinate} {cell.value!r}")
        assert missing == [], f"headers without tooltip: {missing[:10]}"

    def test_column_dictionary_in_run_summary(self, tmp_path):
        wb = self.make(tmp_path)
        vals = [str(r[0]) for r in wb["00_Run_Summary"].iter_rows(values_only=True) if r[0]]
        assert any(v.startswith("COLUMN DICTIONARY") for v in vals)
        assert "07_Welcome_Presence_Pivot" in vals  # dictionary covers new sheets
