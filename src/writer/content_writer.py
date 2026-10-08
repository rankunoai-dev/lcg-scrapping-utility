"""Welcome-section content writer agent.

Pipeline position: AFTER the audit. Inputs per school: the audit observation
(gaps + evidence), the current Welcome text (client CMS), the brand template
library, and Tier-0/approved facts. Output: one draft Welcome section per
school plus a review workbook.

Trust rules (mirroring ARCHITECTURE §2-§3):
  * Facts enter drafts ONLY from: RIO master (name/city/state/street), the
    brand rule (app name), APPROVED locality-register rows, and verbatim
    evidence quoted from the school's own page/current welcome.
  * A template slot with no verified fact becomes a visible
    `[CLIENT INPUT NEEDED: ...]` marker — never invented copy.
  * Sentences of the current Welcome that carry audit-verified PRESENT
    parameters are PRESERVED (healthy content is not disturbed).
  * Deterministic post-guards re-check every draft: preservation survived,
    exact head keywords present, word count in range, and a new-entity scan
    (capitalized phrases not traceable to template/facts/current text are
    flagged). Flagged drafts are marked NEEDS_REVIEW, never silently shipped.

CLI:
    python -m src.writer.content_writer --schools 7440 [--max-cost 0.25]
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
from pathlib import Path
from typing import Optional

from pydantic import BaseModel, Field

from src.evaluator.llm_client import LLMClient, LLMUnavailable
from src.matching.normalize import fold_text, norm_key

logger = logging.getLogger(__name__)

WRITER_VERSION = "1.0.0"
SKILL_PATH = Path("skills/welcome_content_writer/SKILL.md")
TEMPLATE_DIR = Path("config/welcome_templates")
DRAFTS_DIR = Path("data/drafts")

BRAND_TEMPLATE_FILES = {
    "La Petite Academy": "la_petite.txt",
}

PLACEHOLDER_RX = re.compile(r"\[([^\[\]]+)\]")
HEAD_KEYWORDS = {
    "daycare": ["daycare", "day care"],
    "preschool": ["preschool"],
    "child care": ["child care", "childcare"],
    "early education": ["early education"],
}
WORD_RANGE = (150, 340)  # hard guard bounds (target 180-300 lives in the SKILL)


class WriterOutput(BaseModel):
    final_copy: str
    unfilled_placeholders: list[str] = Field(default_factory=list)
    preserved_kept: list[str] = Field(default_factory=list)
    gap_notes: list[str] = Field(default_factory=list)


# ------------------------------------------------------------------ templates
def load_templates(brand: str) -> list[dict]:
    fname = BRAND_TEMPLATE_FILES.get(brand)
    if not fname:
        return []
    path = TEMPLATE_DIR / fname
    if not path.exists():
        return []
    raw = path.read_text(encoding="utf-8")
    templates = []
    for block in re.split(r"^### ", raw, flags=re.M):
        block = block.strip()
        if not block:
            continue
        title, _, body = block.partition("\n")
        body = body.strip()
        templates.append({
            "id": title.strip(),
            "text": body,
            "placeholders": sorted(set(PLACEHOLDER_RX.findall(body))),
        })
    return templates


# ------------------------------------------------------------------ evidence
def _clean_evidence(ev: str | None) -> str | None:
    if not ev:
        return None
    return ev.strip().strip(".").strip().lstrip(".").strip() or None


def _param_quote(obs: dict, pid: str, scopes=("welcome_evidence", "page_evidence"),
                 statuses=("PRESENT", "PARTIAL")) -> str | None:
    p = obs.get("parameters", {}).get(pid, {})
    if p.get("welcome_status") in statuses or p.get("page_status") in statuses:
        for key in scopes:
            ev = _clean_evidence(p.get(key))
            if ev:
                return ev.replace("...", " ").strip()
    return None


# ----------------------------------------------------------------- fact sheet
def build_fact_sheet(school: dict, obs: dict, brand_cfg: dict,
                     register_row: dict | None) -> dict[str, str]:
    """Placeholder -> verified value. Only trusted sources ever enter here."""
    city, state = school.get("city", ""), school.get("state", "")
    facts: dict[str, str] = {
        "City, State": f"{city}, {state}",
        "City": city,
    }
    app = brand_cfg.get("brands", {}).get(school.get("brand", ""), {}).get(
        "app", "SproutAbout")
    safety = _param_quote(obs, "SAFE_01")
    facts["Technology/Safety Feature"] = (
        f"{app}, our live-streaming family app"
        + (f" — and on site: \"{safety}\"" if safety else ""))

    mapping = {
        "Director Name": ("LEAD_01", None),
        "X years serving community": ("LEAD_03", None),
        "X years staff tenure": ("QUAL_04", "LEAD_02"),
        "Landmark 1": ("GEO_06", None),
        "Landmark 2": ("GEO_07", None),
    }
    for slot, (pid, alt) in mapping.items():
        q = _param_quote(obs, pid) or (alt and _param_quote(obs, alt))
        if q:
            facts[slot] = f'from the school\'s own page: "{q}"'

    diff = (_param_quote(obs, "FAC_01") or _param_quote(obs, "THM_07")
            or _param_quote(obs, "QUAL_01"))
    if diff:
        facts["School-Specific Differentiator"] = f'from the school\'s own page: "{diff}"'
    fac = _param_quote(obs, "FAC_01") or _param_quote(obs, "QUAL_03") \
        or _param_quote(obs, "QUAL_01")
    if fac:
        facts["Facility Highlight/Recognition"] = f'from the school\'s own page: "{fac}"'

    if register_row and (register_row.get("approved") or "").strip().lower() == "yes":
        facts["Verified neighborhood (optional to weave in)"] = register_row["locality"]
    return facts


# ------------------------------------------------------------------ preserve
_SENT_RX = re.compile(r"(?<=[.!?])\s+")


def preserve_sentences(current_welcome: str, obs: dict) -> list[str]:
    """Sentences of the current Welcome that carry verified PRESENT evidence."""
    sentences = [s.strip() for s in _SENT_RX.split(fold_text(current_welcome)) if s.strip()]
    cores = []
    for pid, p in obs.get("parameters", {}).items():
        if p.get("welcome_status") == "PRESENT":
            ev = _clean_evidence(p.get("welcome_evidence"))
            if ev:
                cores.append(norm_key(ev.replace("...", " ")))
    keep, seen = [], set()
    for s in sentences:
        skey = norm_key(s)
        if s in seen or not skey:
            continue
        if any(core and (core in skey or skey in core) for core in cores):
            keep.append(s)
            seen.add(s)
    return keep


# ------------------------------------------------------------------ selection
def select_template(templates: list[dict], facts: dict, school_id: str) -> dict | None:
    if not templates:
        return None
    def fillable(t):
        return sum(1 for ph in t["placeholders"] if ph in facts)
    best_score = max(fillable(t) / max(len(t["placeholders"]), 1) for t in templates)
    best = [t for t in templates
            if abs(fillable(t) / max(len(t["placeholders"]), 1) - best_score) < 1e-9]
    return best[int(school_id) % len(best)] if str(school_id).isdigit() else best[0]


# -------------------------------------------------------------------- guards
_CAP_SEQ_RX = re.compile(r"\b[A-Z][A-Za-z'&-]*(?:\s+[A-Z][A-Za-z'&-]*)+\b")
_GUARD_STOPWORDS = {"client input needed", "welcome to our school"}


def run_guards(copy_text: str, preserve: list[str], allowed_corpus: str) -> list[str]:
    flags: list[str] = []
    allowed = norm_key(allowed_corpus)

    for s in preserve:
        skey = norm_key(s)
        out_key = norm_key(copy_text)
        toks = set(skey.split())
        kept = len(toks & set(out_key.split())) / len(toks) if toks else 1.0
        if skey not in out_key and kept < 0.8:
            flags.append(f"PRESERVE_LOST:{s[:60]}")

    low = fold_text(copy_text).lower()
    for kw, variants in HEAD_KEYWORDS.items():
        if not any(v in low for v in variants):
            flags.append(f"MISSING_HEAD_KEYWORD:{kw}")

    wc = len(copy_text.split())
    if not (WORD_RANGE[0] <= wc <= WORD_RANGE[1]):
        flags.append(f"WORD_COUNT_OUT_OF_RANGE:{wc}")

    # scan per sentence so a capitalized sequence never spans a boundary
    # ("...Denny's. Since 1991..." is two sentences, not one entity)
    for sentence in re.split(r"[.!?\n]+", copy_text):
        for m in _CAP_SEQ_RX.finditer(sentence):
            phrase = m.group(0)
            pkey = norm_key(phrase)
            if not pkey or pkey in _GUARD_STOPWORDS:
                continue
            if pkey not in allowed:
                flags.append(f"NEW_ENTITY:{phrase[:50]}")
    return flags


# --------------------------------------------------------------------- agent
class ContentWriter:
    def __init__(self, brand_cfg: dict, llm_client: LLMClient | None = None,
                 skill_path: Path = SKILL_PATH):
        self.brand_cfg = brand_cfg
        self.client = llm_client or LLMClient(
            model=os.environ.get("LCG_WRITER_MODEL", "anthropic/claude-sonnet-5.5"),
            reasoning_effort=os.environ.get("LCG_WRITER_REASONING", "low"))
        self.skill_text = Path(skill_path).read_text(encoding="utf-8")
        DRAFTS_DIR.mkdir(parents=True, exist_ok=True)

    def draft_school(self, school: dict, obs: dict, current_welcome: str,
                     register_row: dict | None = None) -> dict:
        sid = str(school["school_id"])
        templates = load_templates(school.get("brand", ""))
        if not templates:
            return {"school_id": sid, "status": "NO_TEMPLATES",
                    "note": f"No template library for brand "
                            f"'{school.get('brand')}' — add one under "
                            f"{TEMPLATE_DIR}/"}

        facts = build_fact_sheet(school, obs, self.brand_cfg, register_row)
        template = select_template(templates, facts, sid)
        preserve = preserve_sentences(current_welcome, obs)
        gaps = [f"{pid} — {p.get('parameter_name', '')} ({p.get('welcome_status')})"
                for pid, p in obs.get("parameters", {}).items()
                if p.get("welcome_status") in ("ABSENT", "PARTIAL")
                and not pid.startswith("STR_")]

        payload = (
            "## School identity (verified)\n"
            + json.dumps({"school_name": school.get("school_name"),
                          "city": school.get("city"), "state": school.get("state"),
                          "street_address": school.get("street_address")},
                         ensure_ascii=False)
            + "\n\n## TEMPLATE to adapt (" + template["id"] + ")\n" + template["text"]
            + "\n\n## VERIFIED FACT SHEET (the ONLY permitted facts)\n"
            + json.dumps(facts, indent=1, ensure_ascii=False)
            + "\n\n## CURRENT welcome text\n" + (current_welcome or "(none)")
            + "\n\n## PRESERVE these sentences essentially verbatim\n"
            + ("\n".join(f"- {s}" for s in preserve) or "(none)")
            + "\n\n## GAPS to cover\n" + ("\n".join(f"- {g}" for g in gaps) or "(none)")
        )

        try:
            parsed, meta = self.client.evaluate(self.skill_text, payload, WriterOutput)
        except LLMUnavailable as e:
            return {"school_id": sid, "status": "WRITER_FAILED", "note": str(e)}

        allowed_corpus = " ".join([template["text"], current_welcome or "",
                                   json.dumps(facts, ensure_ascii=False),
                                   school.get("school_name", ""), school.get("brand", ""),
                                   school.get("city", ""), school.get("state", ""),
                                   school.get("street_address", "")])
        flags = run_guards(parsed.final_copy, preserve, allowed_corpus)

        draft = {
            "school_id": sid,
            "school_name": school.get("school_name", ""),
            "url": school.get("url", ""),
            "status": "NEEDS_REVIEW" if flags else "CLEAN",
            "template_used": template["id"],
            "draft_copy": parsed.final_copy,
            "client_inputs_needed": parsed.unfilled_placeholders,
            "preserved_sentences": preserve,
            "gap_notes": parsed.gap_notes,
            "guard_flags": flags,
            "word_count": len(parsed.final_copy.split()),
            "writer_model": meta.model,
            "writer_version": WRITER_VERSION,
        }
        (DRAFTS_DIR / f"{sid}_draft.json").write_text(
            json.dumps(draft, indent=1, ensure_ascii=False), encoding="utf-8")
        return draft


# ----------------------------------------------------------------- reporting
def write_drafts_workbook(drafts: list[dict], path="LCG_Welcome_Drafts.xlsx"):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from src.reporter.report_generator import _clean

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Welcome_Drafts"
    headers = ["School ID", "School Name", "School URL", "Status", "Template Used",
               "Word Count", "Draft Copy", "Client Inputs Needed",
               "Preserved Sentences (healthy content kept)", "Gap Coverage Notes",
               "Guard Flags"]
    ws.append(headers)
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for c in ws[1]:
        c.fill = fill
        c.font = Font(bold=True, color="FFFFFF")
    for d in drafts:
        ws.append([_clean(x) for x in (
            d.get("school_id"), d.get("school_name"), d.get("url"),
            d.get("status"), d.get("template_used", ""), d.get("word_count", ""),
            d.get("draft_copy", d.get("note", "")),
            "; ".join(d.get("client_inputs_needed", [])),
            " | ".join(d.get("preserved_sentences", [])),
            " | ".join(d.get("gap_notes", [])),
            "; ".join(d.get("guard_flags", [])))])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=11):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for letter, width in (("A", 10), ("B", 30), ("C", 42), ("D", 14), ("E", 12),
                          ("F", 10), ("G", 95), ("H", 40), ("I", 60), ("J", 45),
                          ("K", 35)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def main():
    import csv
    from dotenv import load_dotenv
    load_dotenv()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

    ap = argparse.ArgumentParser(description="Welcome-section content writer agent")
    ap.add_argument("--schools", required=False, help="comma-separated school ids")
    ap.add_argument("--all-schools", action="store_true")
    ap.add_argument("--max-cost", type=float,
                    default=float(os.environ.get("DEFAULT_MAX_COST_USD", "1.0")))
    ap.add_argument("--output", default="LCG_Welcome_Drafts.xlsx")
    args = ap.parse_args()

    brand_cfg = json.loads(Path("config/brand_defaults.json").read_text(encoding="utf-8"))
    master = {r["school_id"]: r for r in csv.DictReader(
        open("data/master_schools.csv", newline="", encoding="utf-8"))}
    register = {}
    reg_p = Path("data/verified_locality_register.csv")
    if reg_p.exists():
        register = {r["school_id"]: r for r in csv.DictReader(
            open(reg_p, newline="", encoding="utf-8"))}

    wanted = [x.strip() for x in args.schools.split(",")] if args.schools \
        else sorted(master)
    cap = int(os.environ.get("DEFAULT_MAX_SCHOOLS", "10"))
    if len(wanted) > cap and not args.all_schools:
        raise SystemExit(f"GUARDRAIL: {len(wanted)} schools requested; cap is {cap}. "
                         f"Pass --all-schools to run deliberately.")

    writer = ContentWriter(brand_cfg,
                           llm_client=LLMClient(
                               model=os.environ.get("LCG_WRITER_MODEL",
                                                    "anthropic/claude-sonnet-5.5"),
                               max_cost_usd=args.max_cost,
                               reasoning_effort=os.environ.get(
                                   "LCG_WRITER_REASONING", "low")))
    drafts = []
    for sid in wanted:
        school = master.get(sid)
        obs_path = Path(f"data/observations/{sid}_eval.json")
        if school is None or not obs_path.exists():
            logger.warning("[%s] skipped (no master row or no audit observation)", sid)
            continue
        obs = json.loads(obs_path.read_text(encoding="utf-8"))
        # current welcome: client CMS text first (source of truth), scraped fallback
        from src.ingest.welcome_source import ClientWelcomeSource
        current = ClientWelcomeSource().get(sid) or ""
        if not current:
            ex_path = Path(f"data/extracted_sections/{sid}.json")
            if ex_path.exists():
                current = json.loads(ex_path.read_text(encoding="utf-8")).get("welcome_text", "")
        d = writer.draft_school(school, obs, current, register.get(sid))
        logger.info("[%s] %s (template %s, %s words, flags: %s)", sid,
                    d.get("status"), d.get("template_used", "-"),
                    d.get("word_count", "-"), d.get("guard_flags", []))
        drafts.append(d)

    out = write_drafts_workbook(drafts, args.output)
    spent = getattr(writer.client, "total_cost_usd", 0.0)
    logger.info("Drafts: %d | workbook: %s | writer spend: $%.4f",
                len(drafts), out, spent)


if __name__ == "__main__":
    main()
