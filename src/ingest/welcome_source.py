"""Client-attested Welcome content (Tier 0) — WelcomeMessages_Prod_AllSchools.

The client's CMS export is the authoritative source for the WELCOME scope:
it removes scraping/extraction risk entirely for the audit's most important
boundary. Verified 2026-10-05: 0.98-1.00 token overlap with the scraped
welcome on every calibration school (the live pages render this exact text).

Pipeline rules (applied in apply_client_welcome):
  * welcome_text   <- client CMS text when the school is in the file
  * rest scope     <- still the scraped page minus the scraped welcome
                      container (nav/footer excluded) — unchanged
  * fallback       <- scraped welcome when the school is missing from the
                      file (tracked, surfaced in Review Notes)
  * cross-check    <- when both exist, low token overlap flags
                      WELCOME_CMS_PAGE_MISMATCH (live page may be stale)
  * REST_MAY_INCLUDE_WELCOME when the client text exists but the scraper
    found no welcome container to subtract from the rest scope.

CLI (build data/client_welcome.json from the frozen source copy):
    python -m src.ingest.welcome_source
"""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from src.matching.normalize import fold_text, norm_key

SOURCE_XLSX = Path("data/source/WelcomeMessages_Prod_AllSchools_10_2.xlsx")
OUTPUT_JSON = Path("data/client_welcome.json")
REPORT_JSON = Path("data/client_welcome_report.json")

MISMATCH_OVERLAP_THRESHOLD = 0.85


def _norm_sid(value) -> str:
    """'0002' and 2 -> '2' (the file zero-pads school numbers)."""
    s = str(value).strip()
    return str(int(s)) if s.isdigit() else s


def build(source: Path = SOURCE_XLSX, master_csv: Path = Path("data/master_schools.csv")):
    import csv
    import openpyxl

    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    col = {h: i for i, h in enumerate(header)}
    data: dict[str, dict] = {}
    for r in rows:
        if r[col["SchoolNumber"]] is None:
            continue
        sid = _norm_sid(r[col["SchoolNumber"]])
        data[sid] = {
            "brand_code": fold_text(r[col["BrandCode"]]),
            "text": fold_text(r[col["WelcomeMessage"]]),
        }
    wb.close()

    with open(master_csv, newline="", encoding="utf-8") as f:
        master_ids = {row["school_id"] for row in csv.DictReader(f)}

    missing = sorted(master_ids - set(data), key=int)
    extra = sorted(set(data) - master_ids, key=lambda x: int(x) if x.isdigit() else 0)

    OUTPUT_JSON.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source_file": str(source),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "rows": len(data),
        "matched_to_master": len(master_ids & set(data)),
        "master_without_client_welcome": missing,
        "client_rows_not_in_master": extra,
    }
    REPORT_JSON.write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


class ClientWelcomeSource:
    def __init__(self, path: str | Path = OUTPUT_JSON):
        p = Path(path)
        self.data: dict[str, dict] = (
            json.loads(p.read_text(encoding="utf-8")) if p.exists() else {})

    def get(self, school_id) -> str | None:
        entry = self.data.get(_norm_sid(school_id))
        return entry["text"] if entry and entry.get("text") else None


def _token_overlap(a: str, b: str) -> float:
    ta, tb = set(norm_key(a).split()), set(norm_key(b).split())
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def apply_client_welcome(extracted: dict, client_text: str | None) -> dict:
    """Override the welcome scope with client CMS text (in place).

    Returns the extracted dict with welcome_source and source flags set.
    """
    scraped_found = bool(extracted.get("welcome_found"))
    scraped_text = extracted.get("welcome_text", "") or ""
    flags: list[str] = []

    if client_text:
        if scraped_found and scraped_text:
            if _token_overlap(client_text, scraped_text) < MISMATCH_OVERLAP_THRESHOLD:
                flags.append("WELCOME_CMS_PAGE_MISMATCH")
        else:
            flags.append("REST_MAY_INCLUDE_WELCOME")
        extracted["welcome_text"] = client_text
        extracted["welcome_word_count"] = len(client_text.split())
        extracted["welcome_found"] = True
        extracted["extraction_method"] = "CLIENT_FILE"
        extracted["welcome_dom_path"] = "WelcomeMessages_Prod_AllSchools (client CMS export)"
        extracted["welcome_source"] = "CLIENT_FILE"
    else:
        extracted["welcome_source"] = "SCRAPED_FALLBACK" if scraped_found else "NONE"
        if not scraped_found:
            flags.append("NO_WELCOME_ANY_SOURCE")
    extracted["welcome_source_flags"] = flags
    return extracted


if __name__ == "__main__":
    rep = build()
    print(f"client welcome entries: {rep['rows']}")
    print(f"matched to master: {rep['matched_to_master']}/1113")
    print(f"master schools WITHOUT client welcome: {len(rep['master_without_client_welcome'])}"
          f" -> {rep['master_without_client_welcome'][:15]}")
    print(f"client rows not in master: {len(rep['client_rows_not_in_master'])}")
